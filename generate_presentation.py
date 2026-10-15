"""
generate_presentation.py
Generates the comprehensive executive slide deck:
Predictive_Analytics_CA3_Presentation.pptx
for the End-to-End Predictive Analytics Model on Dual Datasets.
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def add_run(p, text):
    """Helper to add run with text in python-pptx."""
    r = p.add_run()
    r.text = text
    return r

def create_deck():
    prs = Presentation()
    # Set 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Definitions
    c_navy = RGBColor(24, 43, 73)        # #182B49 Primary Dark
    c_blue = RGBColor(41, 128, 185)      # #2980B9 Accent Blue
    c_teal = RGBColor(22, 160, 133)      # #16A085 Success/Highlight
    c_dark = RGBColor(44, 62, 80)        # #2C3E50 Body Text
    c_gray = RGBColor(127, 140, 141)     # #7F8C8D Subtitles/Muted
    c_bg_light = RGBColor(245, 247, 250) # #F5F7FA Card Background
    c_white = RGBColor(255, 255, 255)
    c_gold = RGBColor(211, 84, 0)        # Orange/Gold Warning

    def add_header(slide, title_text, category_text="PREDICTIVE ANALYTICS  |  CA3 PRESENTATION"):
        # Header banner shape
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.15))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = c_navy
        top_bar.line.color.rgb = c_navy

        # Category text
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.3))
        p_cat = tb_cat.text_frame.paragraphs[0]
        r_cat = add_run(p_cat, category_text.upper())
        r_cat.font.size = Pt(10)
        r_cat.font.bold = True
        r_cat.font.color.rgb = RGBColor(180, 205, 235)

        # Main Title text
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.7), Inches(0.65))
        p_title = tb_title.text_frame.paragraphs[0]
        r_title = add_run(p_title, title_text)
        r_title.font.size = Pt(20)
        r_title.font.bold = True
        r_title.font.color.rgb = c_white

    def add_card(slide, left, top, width, height, bg_color=c_bg_light, border_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.color.rgb = bg_color
        return card

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = c_navy
    bg1.line.color.rgb = c_navy

    # Inner decorative container
    inner1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.8), Inches(11.333), Inches(5.9))
    inner1.fill.solid()
    inner1.fill.fore_color.rgb = RGBColor(32, 54, 88)
    inner1.line.color.rgb = c_blue
    inner1.line.width = Pt(2)

    tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.3), Inches(10.333), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    r0 = add_run(p0, "PREDICTIVE ANALYTICS  |  END-TO-END MODEL DEVELOPMENT")
    r0.font.size = Pt(12)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(41, 128, 185)
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    r1 = add_run(p1, "Dual-Dataset E-Commerce Consumer Conversion Modeling")
    r1.font.size = Pt(26)
    r1.font.bold = True
    r1.font.color.rgb = c_white
    p1.space_after = Pt(8)

    p2 = tf.add_paragraph()
    r2 = add_run(p2, "Integrating Macro-Session Dynamics (UCI) and Micro-Clickstream Event Funnels (Retailrocket) with Data Leakage Isolation")
    r2.font.size = Pt(14)
    r2.font.color.rgb = RGBColor(200, 215, 235)
    p2.space_after = Pt(28)

    p3 = tf.add_paragraph()
    r3 = add_run(p3, "Candidate / Presenter: Kshitij Bhele\n"
                    "Domain: E-Commerce Predictive Analytics & Web Telemetry Mining\n"
                    "Datasets: UCI Online Shoppers (12,330 sessions) & Retailrocket (2.75M clickstream events)\n"
                    "Academic Presentation: October 2026")
    r3.font.size = Pt(12)
    r3.font.color.rgb = RGBColor(170, 185, 205)

    # =========================================================================
    # SLIDE 2: PROBLEM STATEMENT & MOTIVATION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "1. Problem Statement & Industrial Context")

    # Left Card: The Business Challenge
    add_card(s2, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.6))
    tb_c1 = s2.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.2), Inches(5.4))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True

    p = tf_c1.paragraphs[0]
    r = add_run(p, "The E-Commerce Abandonment Crisis")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(10)

    bullets_c1 = [
        ("The Conversion Deficit: ", "Over 97% of online shopping sessions end without a transaction. Global average conversion rates range between 1.5% and 3.0%."),
        ("Escalating CAC: ", "Customer Acquisition Cost (CAC) via digital ad auctions has risen drastically. Failing to convert acquired visitors destroys operating margins."),
        ("Post-Session Decay: ", "Retrospective retargeting (emails, display remarketing) suffers from severe delay; once a user exits, recovery rates plunge below 3%."),
        ("The Need for Real-Time Intent: ", "E-commerce stores require sub-second machine learning models that evaluate consumer purchase propensity in flight during active browsing.")
    ]
    for b_title, b_desc in bullets_c1:
        p = tf_c1.add_paragraph()
        p.space_after = Pt(8)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11.5)
        r1.font.color.rgb = c_blue
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(11)
        r2.font.color.rgb = c_dark

    # Right Card: The Predictive Analytics Solution
    add_card(s2, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.6))
    tb_c2 = s2.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(5.4))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True

    p = tf_c2.paragraphs[0]
    r = add_run(p, "Predictive Analytics Objective")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(10)

    bullets_c2 = [
        ("In-Flight Intent Scoring: ", "Predict whether an active visitor will convert (Revenue = True / 1) based on early navigational clickstream signals."),
        ("Data Leakage Hazard: ", "Expose and isolate retrospective metric contamination (e.g. Google Analytics' PageValues) that artificially inflates offline test metrics."),
        ("Automated Dynamic Intervention: ", "Trigger automated interventions—such as dynamic 10% vouchers, live chat support, or free-shipping thresholds—before cart abandonment."),
        ("Dual-Dataset Validation: ", "Harmonize macro-session benchmarks (UCI) with high-velocity real-world event logs (Retailrocket) for universal generalizability.")
    ]
    for b_title, b_desc in bullets_c2:
        p = tf_c2.add_paragraph()
        p.space_after = Pt(8)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11.5)
        r1.font.color.rgb = c_teal
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(11)
        r2.font.color.rgb = c_dark

    # =========================================================================
    # SLIDE 3: DUAL DATASET ARCHITECTURE
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "2. Dual-Dataset Architecture: Macro-Session vs. Micro-Clickstream")

    # Card 1: UCI Dataset
    add_card(s3, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.6))
    tb_u = s3.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.2), Inches(5.4))
    tf_u = tb_u.text_frame
    tf_u.word_wrap = True

    p = tf_u.paragraphs[0]
    r = add_run(p, "Dataset 1: UCI Online Shoppers Intention")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(8)

    uci_points = [
        ("Modality: ", "Macro-level session summaries (12,330 visits; 12,205 deduplicated)."),
        ("Target Variable: ", "Revenue (Boolean: 84.5% False, 15.5% True)."),
        ("Behavioral Attributes: ", "Duration & pageview counts on Administrative, Informational, and ProductRelated pages."),
        ("Analytics Telemetry: ", "BounceRates, ExitRates, and the retrospective PageValues metric."),
        ("Contextual Signals: ", "SpecialDay, Month, OperatingSystems, Browser, Region, TrafficType, VisitorType, Weekend."),
        ("Role in Project: ", "Serves as the benchmark for tabular session classification and the empirical testbed for PageValues leakage isolation.")
    ]
    for b_title, b_desc in uci_points:
        p = tf_u.add_paragraph()
        p.space_after = Pt(6)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_blue
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    # Card 2: Retailrocket Dataset
    add_card(s3, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.6))
    tb_r = s3.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(5.4))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    r = add_run(p, "Dataset 2: Retailrocket Clickstream Recommender")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(8)

    rr_points = [
        ("Modality: ", "High-velocity micro-interaction event logs (2,756,101 events across 1,407,580 visitors)."),
        ("Event Taxonomy: ", "view (2,664,312), addtocart (69,332), and transaction (22,457)."),
        ("E-Commerce Funnel: ", "0.81% overall event conversion rate; 0.83% visitor-level purchase rate."),
        ("Enrichment Data: ", "CategoryTree taxonomy (1,669 categories; 25 root departments)."),
        ("Sessionization: ", "Aggregated into visitor journeys with duration, intensity, item depth, and cart ratios."),
        ("Role in Project: ", "Demonstrates high-throughput sessionization, multi-stage funnel progression, and pre-cart vs. full-funnel predictive modeling.")
    ]
    for b_title, b_desc in rr_points:
        p = tf_r.add_paragraph()
        p.space_after = Pt(6)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_teal
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    # =========================================================================
    # SLIDE 4: DATA UNDERSTANDING & EXPLORATORY ANALYSIS
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "3. Data Understanding & Exploratory Data Analysis (EDA)")

    # 3 Stat Cards across top
    stat_boxes = [
        ("2,756,101", "Total Clickstream Events", "Retailrocket event log volume across 4.5 months", c_blue),
        ("0.84%", "View-to-Purchase Funnel", "Extreme conversion attrition across 2.66M views", c_gold),
        ("15.47%", "UCI Conversion Rate", "Class imbalance ratio (1,908 buyers vs 10,422 browsers)", c_teal)
    ]
    for idx, (num, title, sub, color) in enumerate(stat_boxes):
        x = Inches(0.8 + idx * 4.0)
        card = add_card(s4, x, Inches(1.4), Inches(3.7), Inches(1.6), bg_color=c_white, border_color=color)
        tb_s = s4.shapes.add_textbox(x + Inches(0.1), Inches(1.5), Inches(3.5), Inches(1.4))
        tf_s = tb_s.text_frame
        p_n = tf_s.paragraphs[0]
        r_n = add_run(p_n, num)
        r_n.font.size = Pt(24)
        r_n.font.bold = True
        r_n.font.color.rgb = color
        p_t = tf_s.add_paragraph()
        r_t = add_run(p_t, title)
        r_t.font.size = Pt(11)
        r_t.font.bold = True
        r_t.font.color.rgb = c_navy
        p_sub = tf_s.add_paragraph()
        r_sub = add_run(p_sub, sub)
        r_sub.font.size = Pt(9.5)
        r_sub.font.color.rgb = c_gray

    # Bottom Content Card
    add_card(s4, Inches(0.8), Inches(3.2), Inches(11.7), Inches(3.8))
    tb_eda = s4.shapes.add_textbox(Inches(1.0), Inches(3.3), Inches(11.3), Inches(3.6))
    tf_eda = tb_eda.text_frame
    tf_eda.word_wrap = True

    p = tf_eda.paragraphs[0]
    r = add_run(p, "Key Empirical Observations from EDA")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(8)

    eda_insights = [
        ("The Funnel Drop-off Phenomenon: ", "Retailrocket clickstream confirms massive funnel drop-off: out of 100 product views, only 2.60 items are added to cart, and only 0.84 reach transaction completion. Pre-cart intervention is essential."),
        ("Temporal Seasonality: ", "In the UCI dataset, November represents the peak purchasing surge (Cyber Monday / holiday shopping), showing a conversion rate of 25.6%, compared to February with only 1.9%."),
        ("PageValues Bifurcation: ", "Non-converting sessions have a PageValues median of exactly 0.00. Converting sessions exhibit a massive positive spike. This confirms that PageValues is calculated after transactions take place."),
        ("Browsing Velocity vs. Dwell Time: ", "In Retailrocket, converting visitors display higher views_per_item (re-visitation depth: 1.8 vs 1.1) and prolonged session duration, indicating deliberate comparison shopping.")
    ]
    for b_title, b_desc in eda_insights:
        p = tf_eda.add_paragraph()
        p.space_after = Pt(5)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_navy
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    # =========================================================================
    # SLIDE 5: PREPROCESSING & SESSIONIZATION ARCHITECTURE
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "4. Preprocessing & Clickstream Sessionization Architecture")

    # Left: UCI Preprocessing Pipeline
    add_card(s5, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.6))
    tb_p1 = s5.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.2), Inches(5.4))
    tf_p1 = tb_p1.text_frame
    tf_p1.word_wrap = True

    p = tf_p1.paragraphs[0]
    r = add_run(p, "UCI Data Preprocessing Pipeline")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(8)

    p1_steps = [
        ("Deduplication: ", "125 exact duplicate sessions removed (12,205 unique sessions retained)."),
        ("Missing Value Checks: ", "0 nulls identified across all 18 features; full integrity verified."),
        ("StandardScaler: ", "14 numerical columns scaled to mean 0, unit variance (durations, bounce/exit rates)."),
        ("OneHotEncoder: ", "Month, VisitorType, and Weekend encoded with handle_unknown='ignore'."),
        ("ColumnTransformer: ", "Scikit-Learn preprocessor eliminates data leakage across cross-validation folds."),
        ("Stratified Split: ", "80/20 train/test partition maintaining 15.5% conversion ratio.")
    ]
    for b_title, b_desc in p1_steps:
        p = tf_p1.add_paragraph()
        p.space_after = Pt(6)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_blue
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    # Right: Retailrocket Clickstream Sessionization
    add_card(s5, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.6))
    tb_p2 = s5.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(5.4))
    tf_p2 = tb_p2.text_frame
    tf_p2.word_wrap = True

    p = tf_p2.paragraphs[0]
    r = add_run(p, "Retailrocket Sessionization Pipeline")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(8)

    p2_steps = [
        ("Vectorised Clickstream Grouping: ", "Aggregated 2.75M event logs by visitorid in 1.5 seconds."),
        ("Behavioral Feature Engineering: ", "Engineered duration_min, view_intensity, unique_items, views_per_item, and cart_to_view_ratio."),
        ("Funnel Target Construction: ", "has_transaction (1 if completed transaction, 0 if abandoned)."),
        ("Stratified Modeling Sample: ", "Combined all 11,719 buyers + 40,000 non-buyers (51,719 visitors total; 22.7% positive class)."),
        ("Early Pre-Cart Subsetting: ", "Strict pre-cart feature set: views, unique_items, duration_min, view_intensity, views_per_item."),
        ("Full-Funnel Journey Subsetting: ", "Adds late-stage signals: cart_adds and cart_to_view_ratio.")
    ]
    for b_title, b_desc in p2_steps:
        p = tf_p2.add_paragraph()
        p.space_after = Pt(6)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_teal
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    # =========================================================================
    # SLIDE 6: THE FEATURE LEAKAGE EXPERIMENT
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "5. The Feature Leakage Experiment: The PageValues Anomaly")

    # Warning Box
    add_card(s6, Inches(0.8), Inches(1.4), Inches(11.7), Inches(1.5), bg_color=RGBColor(253, 237, 236), border_color=RGBColor(231, 76, 60))
    tb_w = s6.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(1.3))
    tf_w = tb_w.text_frame
    tf_w.word_wrap = True
    p = tf_w.paragraphs[0]
    r = add_run(p, "CRITICAL INSIGHT: RETROSPECTIVE DATA LEAKAGE IN E-COMMERCE ANALYTICS")
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(192, 57, 43)
    p = tf_w.add_paragraph()
    r = add_run(p, "Google Analytics defines PageValue = (Ecommerce Revenue + Goal Value) / Unique Pageviews. "
                   "Because PageValues incorporates the revenue realized at the conclusion of the session, including it creates severe retrospective leakage. "
                   "Models achieve inflated ROC-AUC (>0.93) in offline validation, but catastrophically fail in real-time production.")
    r.font.size = Pt(11)
    r.font.color.rgb = c_dark

    # Bottom Two Comparison Cards
    add_card(s6, Inches(0.8), Inches(3.1), Inches(5.6), Inches(3.9))
    tb_l1 = s6.shapes.add_textbox(Inches(1.0), Inches(3.2), Inches(5.2), Inches(3.7))
    tf_l1 = tb_l1.text_frame
    tf_l1.word_wrap = True
    p = tf_l1.paragraphs[0]
    r = add_run(p, "Configuration A: With PageValues (Retrospective)")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(192, 57, 43)
    l1_text = [
        ("Metric Distortion: ", "Gradient Boosting reaches 0.9372 ROC-AUC; Decision Tree Recall reaches 85.86%."),
        ("Dominant Feature: ", "PageValues absorbs over 83.4% of total feature importance."),
        ("Production Reality: ", "Completely useless in a live production environment because transactions have not yet occurred!"),
        ("Model Trap: ", "Creates a false sense of model confidence while obscuring genuine behavioral exploration patterns.")
    ]
    for b_title, b_desc in l1_text:
        p = tf_l1.add_paragraph()
        p.space_after = Pt(4)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = RGBColor(192, 57, 43)
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    add_card(s6, Inches(6.8), Inches(3.1), Inches(5.7), Inches(3.9))
    tb_l2 = s6.shapes.add_textbox(Inches(7.0), Inches(3.2), Inches(5.3), Inches(3.7))
    tf_l2 = tb_l2.text_frame
    tf_l2.word_wrap = True
    p = tf_l2.paragraphs[0]
    r = add_run(p, "Configuration B: Without PageValues (Production-Safe)")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = c_teal
    l2_text = [
        ("Honest Evaluation: ", "Gradient Boosting achieves 0.7843 ROC-AUC; Logistic Regression achieves 0.7509 ROC-AUC."),
        ("Production Champion: ", "Balanced Logistic Regression achieves 76.70% Recall, successfully capturing 3 out of 4 converting shoppers early."),
        ("True Behavioral Signals: ", "Model relies on ProductRelated_Duration (31.2%), ExitRates (18.5%), and Month_Nov (7.4%)."),
        ("Live Deployment: ", "100% operational in real time; enables dynamic interventions before visitor leaves.")
    ]
    for b_title, b_desc in l2_text:
        p = tf_l2.add_paragraph()
        p.space_after = Pt(4)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = c_teal
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    # =========================================================================
    # SLIDE 7: PREDICTIVE MODELLING APPROACH
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "6. Predictive Modelling Approach: 5 Machine Learning Algorithms")

    algos = [
        ("Logistic Regression", "Probabilistic Linear Classifier", "Serves as an interpretable parametric baseline. Implements cost-sensitive class weighting to counteract severe class imbalance.", c_blue),
        ("Decision Tree (CART)", "Non-Linear Rule-Based Classifier", "Recursively partitions the feature space based on Gini impurity. Identifies non-linear threshold interactions without requiring linearity.", c_navy),
        ("Random Forest", "Ensemble Bagging Architecture", "Averages predictions across 50-100 decorrelated decision trees built on bootstrap samples. Substantially reduces model variance.", c_teal),
        ("Gradient Tree Boosting", "Ensemble Sequential Boosting", "Sequentially minimizes binomial log-loss by fitting shallow regression trees to pseudo-residuals. Maximizes discriminative power.", c_gold),
        ("Dummy Baseline", "Stratified Random Guessing", "Generates stratified predictions matching empirical training class frequencies. Establishes the performance floor (ROC-AUC ~ 0.50).", c_gray)
    ]
    for idx, (name, role, desc, color) in enumerate(algos):
        y = Inches(1.4 + idx * 1.15)
        add_card(s7, Inches(0.8), y, Inches(11.7), Inches(1.0), bg_color=c_white, border_color=color)
        tb_a = s7.shapes.add_textbox(Inches(1.0), y + Inches(0.08), Inches(11.3), Inches(0.85))
        tf_a = tb_a.text_frame
        p = tf_a.paragraphs[0]
        r1 = add_run(p, f"{idx+1}. {name}  ")
        r1.font.size = Pt(13)
        r1.font.bold = True
        r1.font.color.rgb = color
        r2 = add_run(p, f"[{role}]")
        r2.font.size = Pt(10.5)
        r2.font.italic = True
        r2.font.color.rgb = c_gray
        p_d = tf_a.add_paragraph()
        r_d = add_run(p_d, desc)
        r_d.font.size = Pt(10)
        r_d.font.color.rgb = c_dark

    # =========================================================================
    # SLIDE 8: MODEL TRAINING & HYPERPARAMETER TUNING
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "7. Model Training & Hyperparameter Tuning (Stratified 5-Fold CV)")

    # Left Card: Protocol
    add_card(s8, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.6))
    tb_cv1 = s8.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.2), Inches(5.4))
    tf_cv1 = tb_cv1.text_frame
    tf_cv1.word_wrap = True

    p = tf_cv1.paragraphs[0]
    r = add_run(p, "Cross-Validation Protocol")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(8)

    cv_points = [
        ("Stratified 5-Fold Partitioning: ", "Preserves exact class ratios (15.5% in UCI, 22.7% in Retailrocket sample) within each training and validation fold."),
        ("Optimization Metric (ROC-AUC): ", "Hyperparameters tuned via GridSearchCV optimizing for Area Under ROC Curve rather than misleading accuracy."),
        ("Leakage Elimination: ", "Feature scaling and encoding transformations fit strictly within training folds, preventing validation data contamination."),
        ("Reproducibility: ", "Random state fixed (random_state=42) across all splits, estimators, and cross-validation folds.")
    ]
    for b_title, b_desc in cv_points:
        p = tf_cv1.add_paragraph()
        p.space_after = Pt(8)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_blue
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    # Right Card: Tuning Grid Specifications
    add_card(s8, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.6))
    tb_cv2 = s8.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(5.4))
    tf_cv2 = tb_cv2.text_frame
    tf_cv2.word_wrap = True

    p = tf_cv2.paragraphs[0]
    r = add_run(p, "Hyperparameter Search Grids")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(8)

    grid_specs = [
        ("Logistic Regression: ", "Inverse regularization C in {0.1, 1.0}; solver='lbfgs'; max_iter=1000; class_weight='balanced'."),
        ("Decision Tree: ", "max_depth in {4, 7}; class_weight='balanced'; criterion='gini' to avoid deep overfitted leaves."),
        ("Random Forest: ", "n_estimators in {50, 100}; max_depth in {6, 10}; class_weight='balanced'; n_jobs=-1."),
        ("Gradient Boosting: ", "n_estimators in {50, 100}; learning_rate in {0.05, 0.1}; max_depth=3-4 sequential trees.")
    ]
    for b_title, b_desc in grid_specs:
        p = tf_cv2.add_paragraph()
        p.space_after = Pt(8)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_teal
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    # =========================================================================
    # SLIDE 9: MODEL EVALUATION & BENCHMARK PERFORMANCE
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "8. Model Evaluation & Benchmark Performance Across Configurations")

    # Benchmark Table
    tbl_shape = s9.shapes.add_table(rows=9, cols=7, left=Inches(0.8), top=Inches(1.4), width=Inches(11.7), height=Inches(4.2))
    tbl = tbl_shape.table

    headers = ["Dataset Setup", "Top Performing Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"]
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = c_navy
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = c_white

    table_data = [
        ("UCI (With PageValues)", "Gradient Boosting", "0.9062", "0.7284", "0.6387", "0.6806", "0.9372"),
        ("UCI (With PageValues)", "Logistic Regression", "0.8525", "0.5190", "0.7880", "0.6258", "0.9109"),
        ("UCI (Without PageValues - Safe)", "Gradient Boosting (High Precision)", "0.8415", "0.4528", "0.0628", "0.1103", "0.7843"),
        ("UCI (Without PageValues - Safe)", "Logistic Regression (Production Champion)", "0.6346", "0.2673", "0.7670", "0.3965", "0.7509"),
        ("Retailrocket (Pre-Cart Intent)", "Gradient Boosting (Best F1)", "0.9030", "0.7837", "0.7901", "0.7869", "0.9630"),
        ("Retailrocket (Pre-Cart Intent)", "Random Forest (High Recall)", "0.8728", "0.6474", "0.9633", "0.7743", "0.9619"),
        ("Retailrocket (Full-Funnel Journey)", "Gradient Boosting", "0.9757", "0.9544", "0.9377", "0.9460", "0.9915"),
        ("Baseline (Both Datasets)", "Dummy Classifier", "0.6588 - 0.7374", "0.1636 - 0.2386", "0.1649 - 0.2308", "0.1643 - 0.2347", "0.5043 - 0.5075")
    ]
    for row_idx, r_data in enumerate(table_data, start=1):
        for col_idx, val in enumerate(r_data):
            cell = tbl.cell(row_idx, col_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(9.5)
            if col_idx in [1, 6]:
                p.font.bold = True
            cell.fill.solid()
            if row_idx in [1, 2]:
                cell.fill.fore_color.rgb = RGBColor(235, 245, 251)
            elif row_idx in [3, 4]:
                cell.fill.fore_color.rgb = RGBColor(254, 249, 231)
            elif row_idx in [5, 6]:
                cell.fill.fore_color.rgb = RGBColor(234, 250, 241)
            elif row_idx == 7:
                cell.fill.fore_color.rgb = RGBColor(244, 236, 247)
            else:
                cell.fill.fore_color.rgb = RGBColor(245, 245, 245)

    # Highlight Callout
    add_card(s9, Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.3), bg_color=c_white, border_color=c_teal)
    tb_k = s9.shapes.add_textbox(Inches(1.0), Inches(5.9), Inches(11.3), Inches(1.1))
    tf_k = tb_k.text_frame
    p = tf_k.paragraphs[0]
    r = add_run(p, "KEY EVALUATION TAKEAWAYS:\n")
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = c_navy
    r2 = add_run(p, "1. In UCI without PageValues, Balanced Logistic Regression captures 76.70% of buyers, acting as the ideal early-warning trigger.\n"
                    "2. In Retailrocket, Pre-Cart Gradient Boosting achieves 0.9630 ROC-AUC & 79.01% Recall using only browsing telemetry before any cart additions occur.")
    r2.font.size = Pt(10)
    r2.font.color.rgb = c_dark

    # =========================================================================
    # SLIDE 10: MODEL INTERPRETATION & FEATURE IMPORTANCE
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "9. Model Interpretation: What Drives E-Commerce Conversion?")

    # Left: UCI Feature Driver Card
    add_card(s10, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.6))
    tb_i1 = s10.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.2), Inches(5.4))
    tf_i1 = tb_i1.text_frame
    tf_i1.word_wrap = True

    p = tf_i1.paragraphs[0]
    r = add_run(p, "UCI Production-Safe Feature Drivers")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(8)

    uci_drivers = [
        ("ProductRelated_Duration (31.2%): ", "Total active time spent browsing product pages is the #1 genuine predictor of purchase intention."),
        ("ExitRates (18.5%): ", "Low exit rates strongly correlate with high buying intent; bouncing or rapidly exiting visitors rarely convert."),
        ("ProductRelated Pageviews (14.7%): ", "Exploration breadth across multiple SKUs signifies active product evaluation."),
        ("Administrative Duration (8.9%): ", "Time spent checking shipping policies, account details, or terms signifies intent to finalize order."),
        ("Month_Nov (7.4%): ", "Seasonal holiday promotional peak substantially lifts visitor baseline propensity.")
    ]
    for b_title, b_desc in uci_drivers:
        p = tf_i1.add_paragraph()
        p.space_after = Pt(6)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_blue
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    # Right: Retailrocket Feature Driver Card
    add_card(s10, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.6))
    tb_i2 = s10.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(5.4))
    tf_i2 = tb_i2.text_frame
    tf_i2.word_wrap = True

    p = tf_i2.paragraphs[0]
    r = add_run(p, "Retailrocket Pre-Cart Clickstream Drivers")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(8)

    rr_drivers = [
        ("views_per_item (38.6%): ", "Re-visitation depth: repeatedly returning to the same item indicates deliberative decision-making."),
        ("duration_min (27.4%): ", "Active session duration: prolonged exploration strongly differentiates buyers from transient bounces."),
        ("view_intensity (19.8%): ", "Browsing velocity: rapid page navigation per minute indicates urgent consumer evaluation."),
        ("unique_items (9.2%): ", "Targeted vs broad catalog exploration: buyers focus on a tight consideration set."),
        ("cart_to_view_ratio (Full-Funnel): ", "When cart actions are included, cart additions dominate the funnel with >60% split importance.")
    ]
    for b_title, b_desc in rr_drivers:
        p = tf_i2.add_paragraph()
        p.space_after = Pt(6)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_teal
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    # =========================================================================
    # SLIDE 11: REAL-TIME INFERENCE & INTERVENTION ENGINE
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "10. Real-Time Inference & Automated Dynamic Intervention Engine")

    # Workflow Box across Top
    add_card(s11, Inches(0.8), Inches(1.4), Inches(11.7), Inches(1.3), bg_color=c_navy)
    tb_eng = s11.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(1.1))
    tf_eng = tb_eng.text_frame
    p = tf_eng.paragraphs[0]
    r = add_run(p, "REAL-TIME PREDICTIVE DECISION ARCHITECTURE")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(180, 205, 235)
    p = tf_eng.add_paragraph()
    r = add_run(p, "Clickstream Telemetry Stream  -->  Feature Scaling (<5ms)  -->  Model Inference (<10ms)  -->  Threshold Engine  -->  Dynamic Business Action")
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = c_white

    # Case 1 & Case 2
    add_card(s11, Inches(0.8), Inches(2.9), Inches(5.6), Inches(4.1))
    tb_sim1 = s11.shapes.add_textbox(Inches(1.0), Inches(3.0), Inches(5.2), Inches(3.9))
    tf_sim1 = tb_sim1.text_frame
    tf_sim1.word_wrap = True

    p = tf_sim1.paragraphs[0]
    r = add_run(p, "Live Case 1: Macro Session Visitor (UCI)")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = c_navy
    sim1_lines = [
        ("Visitor Session Profile: ", "30 Product Pages, 980s active time, Nov session, New Visitor."),
        ("Leakage-Free Score: ", "Predicted Purchase Probability = 38.87%."),
        ("Intervention Threshold: ", "Crosses High-Intent Threshold (P >= 35%)."),
        ("Dynamic Action Triggered: ", "Deploy timed 10% instant checkout voucher overlay to secure transaction before tab closure."),
        ("Business ROI: ", "Converts a high-intent fence-sitter into confirmed revenue.")
    ]
    for b_title, b_desc in sim1_lines:
        p = tf_sim1.add_paragraph()
        p.space_after = Pt(4)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = c_blue
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    add_card(s11, Inches(6.8), Inches(2.9), Inches(5.7), Inches(4.1))
    tb_sim2 = s11.shapes.add_textbox(Inches(7.0), Inches(3.0), Inches(5.3), Inches(3.9))
    tf_sim2 = tb_sim2.text_frame
    tf_sim2.word_wrap = True

    p = tf_sim2.paragraphs[0]
    r = add_run(p, "Live Case 2: Clickstream Visitor (Retailrocket)")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = c_navy
    sim2_lines = [
        ("Clickstream Profile: ", "12 Product Views, 7 Unique Items, 8.5 min dwell time (pre-cart)."),
        ("Pre-Cart Score: ", "Predicted Purchase Probability = 20.21%."),
        ("Intervention Threshold: ", "Exploratory Consideration Stage (15% <= P < 30%)."),
        ("Dynamic Action Triggered: ", "Display related trending cross-category items & free shipping threshold progress bar."),
        ("Business ROI: ", "Nudges browsing visitor toward adding their first item to cart.")
    ]
    for b_title, b_desc in sim2_lines:
        p = tf_sim2.add_paragraph()
        p.space_after = Pt(4)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = c_teal
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    # =========================================================================
    # SLIDE 12: CONCLUSION & FUTURE DIRECTIONS
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "11. Conclusion, Takeaways & Future Directions")

    add_card(s12, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.6))
    tb_c12 = s12.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.2), Inches(5.4))
    tf_c12 = tb_c12.text_frame
    tf_c12.word_wrap = True

    p = tf_c12.paragraphs[0]
    r = add_run(p, "Summary of Project Contributions")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(8)

    c12_points = [
        ("Rigorous KDD Implementation: ", "Demonstrated all 5 KDD stages on two real-world datasets spanning 2.75M interactions."),
        ("Data Leakage Resolved: ", "Proven that retrospective PageValues distorts evaluation; built genuine leakage-free pipelines."),
        ("Benchmark Performance: ", "Balanced Logistic Regression captures 76.7% of UCI buyers; Retailrocket Pre-Cart model achieves 0.963 ROC-AUC."),
        ("Business Relevance: ", "Built an automated real-time intervention engine for dynamic checkout incentives.")
    ]
    for b_title, b_desc in c12_points:
        p = tf_c12.add_paragraph()
        p.space_after = Pt(8)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_blue
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    add_card(s12, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.6))
    tb_f12 = s12.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(5.4))
    tf_f12 = tb_f12.text_frame
    tf_f12.word_wrap = True

    p = tf_f12.paragraphs[0]
    r = add_run(p, "Limitations & Future Work")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = c_navy
    p.space_after = Pt(8)

    f12_points = [
        ("Cross-Device Stitching: ", "Current telemetry cannot link mobile exploration with desktop checkout; future work should integrate graph identity resolution."),
        ("Deep Sequential Learning: ", "Explore Transformer-based session architectures (e.g. BERT4Rec) to model item-to-item transition embeddings."),
        ("Reinforcement Learning: ", "Deploy an online Multi-Armed Bandit (MAB) to learn optimal real-time dynamic discount values without margin erosion."),
        ("Multi-Modal Integration: ", "Incorporate product thumbnail image embeddings and customer review sentiment into the clickstream model.")
    ]
    for b_title, b_desc in f12_points:
        p = tf_f12.add_paragraph()
        p.space_after = Pt(8)
        r1 = add_run(p, "• " + b_title)
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_teal
        r2 = add_run(p, b_desc)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_dark

    prs.save("Predictive_Analytics_CA3_Presentation.pptx")
    print("[+] Successfully generated Predictive_Analytics_CA3_Presentation.pptx!")

if __name__ == "__main__":
    create_deck()
